"""
RAG Data Ingestion Pipeline
Handles ingestion of FAQs, docs, product manuals, and other knowledge sources
"""

import os
import json
import logging
from typing import List, Dict, Any, Optional
from pathlib import Path
import markdown
from bs4 import BeautifulSoup
try:
    import docx
except ImportError:
    docx = None
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class RAGDataIngestion:
    """Data ingestion pipeline for RAG system"""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", " ", ""]
        )
        
    def ingest_faq_database(self, faq_data: List[Dict]) -> List[Document]:
        """Ingest FAQ database into LangChain documents"""
        documents = []
        
        for faq in faq_data:
            # Create comprehensive text for each FAQ
            content = f"""
            Question: {faq['question']}
            Answer: {faq['answer']}
            Category: {faq['category']}
            Keywords: {', '.join(faq['keywords'])}
            """
            
            # Create metadata
            metadata = {
                'source': 'faq_database',
                'category': faq['category'],
                'type': 'faq',
                'question': faq['question'],
                'keywords': faq['keywords']
            }
            
            # Create LangChain document
            doc = Document(page_content=content.strip(), metadata=metadata)
            documents.append(doc)
            
        logger.info(f"Ingested {len(documents)} FAQ entries")
        return documents
    
    def ingest_text_file(self, file_path: str, source_type: str = "text") -> List[Document]:
        """Ingest a text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()
            
            # Split into chunks
            chunks = self.text_splitter.split_text(content)
            
            documents = []
            for i, chunk in enumerate(chunks):
                metadata = {
                    'source': file_path,
                    'type': source_type,
                    'chunk_id': i,
                    'total_chunks': len(chunks)
                }
                doc = Document(page_content=chunk, metadata=metadata)
                documents.append(doc)
            
            logger.info(f"Ingested {len(documents)} chunks from {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Error ingesting {file_path}: {e}")
            return []
    
    def ingest_pdf_file(self, file_path: str) -> List[Document]:
        """Ingest a PDF file"""
        try:
            if PdfReader is None:
                logger.warning("pypdf not available, skipping PDF file")
                return []
                
            reader = PdfReader(file_path)
            text = ""
            
            for page in reader.pages:
                text += page.extract_text() + "\n"
            
            # Split into chunks
            chunks = self.text_splitter.split_text(text)
            
            documents = []
            for i, chunk in enumerate(chunks):
                metadata = {
                    'source': file_path,
                    'type': 'pdf',
                    'chunk_id': i,
                    'total_chunks': len(chunks)
                }
                doc = Document(page_content=chunk, metadata=metadata)
                documents.append(doc)
            
            logger.info(f"Ingested {len(documents)} chunks from PDF {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Error ingesting PDF {file_path}: {e}")
            return []
    
    def ingest_docx_file(self, file_path: str) -> List[Document]:
        """Ingest a DOCX file"""
        try:
            if docx is None:
                logger.warning("python-docx not available, skipping DOCX file")
                return []
                
            doc = docx.Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Split into chunks
            chunks = self.text_splitter.split_text(text)
            
            documents = []
            for i, chunk in enumerate(chunks):
                metadata = {
                    'source': file_path,
                    'type': 'docx',
                    'chunk_id': i,
                    'total_chunks': len(chunks)
                }
                doc = Document(page_content=chunk, metadata=metadata)
                documents.append(doc)
            
            logger.info(f"Ingested {len(documents)} chunks from DOCX {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Error ingesting DOCX {file_path}: {e}")
            return []
    
    def ingest_markdown_file(self, file_path: str) -> List[Document]:
        """Ingest a Markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()
            
            # Convert markdown to HTML then extract text
            html = markdown.markdown(content)
            soup = BeautifulSoup(html, 'html.parser')
            text = soup.get_text()
            
            # Split into chunks
            chunks = self.text_splitter.split_text(text)
            
            documents = []
            for i, chunk in enumerate(chunks):
                metadata = {
                    'source': file_path,
                    'type': 'markdown',
                    'chunk_id': i,
                    'total_chunks': len(chunks)
                }
                doc = Document(page_content=chunk, metadata=metadata)
                documents.append(doc)
            
            logger.info(f"Ingested {len(documents)} chunks from Markdown {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Error ingesting Markdown {file_path}: {e}")
            return []
    
    def ingest_directory(self, directory_path: str) -> List[Document]:
        """Ingest all supported files from a directory"""
        documents = []
        directory = Path(directory_path)
        
        if not directory.exists():
            logger.error(f"Directory {directory_path} does not exist")
            return documents
        
        # Supported file extensions
        supported_extensions = {
            '.txt': self.ingest_text_file,
            '.md': self.ingest_markdown_file,
            '.pdf': self.ingest_pdf_file,
            '.docx': self.ingest_docx_file
        }
        
        for file_path in directory.rglob('*'):
            if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
                file_docs = supported_extensions[file_path.suffix.lower()](str(file_path))
                documents.extend(file_docs)
        
        logger.info(f"Ingested {len(documents)} total documents from {directory_path}")
        return documents
    
    def ingest_product_manuals(self, manuals_data: List[Dict]) -> List[Document]:
        """Ingest product manual data"""
        documents = []
        
        for manual in manuals_data:
            content = f"""
            Product: {manual.get('product_name', 'Unknown')}
            Model: {manual.get('model', 'Unknown')}
            Category: {manual.get('category', 'Unknown')}
            Content: {manual.get('content', '')}
            Features: {', '.join(manual.get('features', []))}
            Specifications: {manual.get('specifications', '')}
            Troubleshooting: {manual.get('troubleshooting', '')}
            """
            
            metadata = {
                'source': 'product_manuals',
                'type': 'product_manual',
                'product_name': manual.get('product_name', 'Unknown'),
                'model': manual.get('model', 'Unknown'),
                'category': manual.get('category', 'Unknown')
            }
            
            doc = Document(page_content=content.strip(), metadata=metadata)
            documents.append(doc)
        
        logger.info(f"Ingested {len(documents)} product manuals")
        return documents
    
    def create_sample_knowledge_base(self) -> List[Document]:
        """Create a sample knowledge base with various document types"""
        documents = []
        
        # Sample product manuals
        sample_manuals = [
            {
                'product_name': 'TechyMart Wireless Headphones Pro',
                'model': 'TM-WH-2024',
                'category': 'audio',
                'content': 'Premium wireless headphones with noise cancellation, 30-hour battery life, and premium sound quality.',
                'features': ['Active Noise Cancellation', '30-hour battery', 'Quick Charge', 'Premium Sound'],
                'specifications': 'Driver: 40mm, Frequency: 20Hz-20kHz, Battery: 30h, Weight: 250g',
                'troubleshooting': 'If headphones won\'t connect: 1) Reset by holding power button for 10 seconds 2) Clear Bluetooth cache 3) Contact support'
            },
            {
                'product_name': 'TechyMart Gaming Mouse Elite',
                'model': 'TM-GM-2024',
                'category': 'gaming',
                'content': 'High-precision gaming mouse with customizable RGB lighting and programmable buttons.',
                'features': ['16000 DPI', 'RGB Lighting', 'Programmable Buttons', 'Ergonomic Design'],
                'specifications': 'DPI: 16000, Polling Rate: 1000Hz, Buttons: 8, Weight: 100g',
                'troubleshooting': 'If mouse is unresponsive: 1) Check USB connection 2) Update drivers 3) Try different USB port'
            }
        ]
        
        documents.extend(self.ingest_product_manuals(sample_manuals))
        
        # Sample policy documents
        policy_content = """
        TechyMart Return Policy
        
        1. 30-Day Return Window: All products can be returned within 30 days of purchase.
        2. Condition Requirements: Items must be in original condition with all packaging and accessories.
        3. Electronics: Must include original packaging and all included components.
        4. Return Process: Start return online or contact customer service for assistance.
        5. Refund Timeline: Refunds processed within 3-5 business days after receiving returned item.
        6. Return Shipping: Free return shipping provided for defective items.
        7. Exceptions: Custom items and personalized products are not returnable.
        """
        
        policy_doc = Document(
            page_content=policy_content,
            metadata={'source': 'policy_documents', 'type': 'return_policy', 'title': 'Return Policy'}
        )
        documents.append(policy_doc)
        
        # Sample technical support guide
        tech_support_content = """
        TechyMart Technical Support Guide
        
        Common Issues and Solutions:
        
        1. Product Not Working:
           - Check power connections
           - Verify all cables are properly connected
           - Try different power outlet
           - Contact technical support if issue persists
        
        2. Software Installation Issues:
           - Ensure system meets minimum requirements
           - Download latest drivers from manufacturer website
           - Run installation as administrator
           - Check antivirus software settings
        
        3. Connectivity Problems:
           - Restart router and device
           - Check network settings
           - Update device firmware
           - Contact ISP if internet issues persist
        
        4. Performance Issues:
           - Close unnecessary applications
           - Check available storage space
           - Update device drivers
           - Run system diagnostics
        """
        
        tech_doc = Document(
            page_content=tech_support_content,
            metadata={'source': 'technical_guides', 'type': 'troubleshooting', 'title': 'Technical Support Guide'}
        )
        documents.append(tech_doc)
        
        logger.info(f"Created sample knowledge base with {len(documents)} documents")
        return documents

# Global data ingestion instance
rag_data_ingestion = RAGDataIngestion()
